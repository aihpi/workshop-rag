"""Build the claim slides of RAG III (Mentimeter rounds, definitions, breakdowns) from the AISC Light template.

    uv run --with python-pptx python slides/src/build_claims.py de slides/src/rag_iii_claims_german.pptx
    uv run --with python-pptx python slides/src/build_claims.py en slides/src/rag_iii_claims_english.pptx
"""

import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TEMPLATE = Path('/home/hanmul/.claude/plugins/synced/19464ea5-effb-41ee-92e0-029ec34f55ca_126a6649-ed73-4cb3-a3d3-fbe1ee1e4ac5/'
                'kisz-house-style/skills/hpi-corporate-design/assets/templates/202608_AISC-Template_Light.pptx')
INK, GREY, RULE, ACCENT, NOTE = RGBColor(0x21, 0x24, 0x27), RGBColor(0x5a, 0x60, 0x65), RGBColor(0xd6, 0xd7, 0xd8), RGBColor(0xb1, 0x06, 0x3a), RGBColor(0xfd, 0xe9, 0xbf)

T = {
    'de': {
        'menti1': 'Errate die Zahl',
        'menti2': 'Errate die Zahl, Runde 2',
        'code': 'Mentimeter-Code: ____',
        'text': 'Text',
        'defs_title': 'Was ist eine Aussage?',
        'defs': [
            ('Alltag (Wörterbuch)', 'Eine Aussage ist ein Satz, der etwas feststellt.',
             'Danach ist der Beispielsatz eine Aussage. Man kann auch zwei zählen: Hauptsatz und Nebensatz sind zwei Sätze.'),
            ('Logik', 'Eine Aussage ist etwas, das für sich allein wahr oder falsch sein kann.',
             '"A, weil B und C" sind drei Aussagen: A, B, C, und außerdem die Behauptung, dass B und C der Grund für A sind. Danach hat der Beispielsatz drei bis vier Aussagen.'),
            ('NLP-Forschung (FActScore)', 'Eine Aussage ist die kleinste Einheit, die eine Information trägt und sich gegen eine Quelle prüfen lässt.',
             'Jedes Element einer Aufzählung zählt. Danach hat der Beispielsatz drei Aussagen.'),
        ],
        'work_title': 'Arbeitsdefinition',
        'work': 'Eine Aussage ist der kleinste Teil der Referenzantwort, der sich für sich allein gegen das Kompendium prüfen lässt.',
        'work_rules': ['Jedes Element einer Aufzählung ist eine eigene Aussage.',
                       'Jede Begründung (weil, da) und jeder Zweck (um zu) ist eine eigene Aussage.',
                       'Wiederholungen und Füllwörter zählen nicht.'],
        'break_title': 'Beispiel {n}: Wie viele Aussagen?',
        'rows': ['Alltag', 'Logik', 'Arbeitsdefinition'],
        'ragas_title': 'Was RAGAS dem Richter sagt',
        'ragas_lead': 'Die Anweisung an den Richter, wörtlich:',
        'ragas_none': 'Keine Definition. Der Richter entscheidet selbst, was eine Aussage ist.',
        'ragas_cols': [('gpt-oss-120b', '2 Aussagen', 'zerlegt meist je Satz'),
                       ('llama-3-3-70b', '4 Aussagen', 'zerlegt feiner')],
        'ragas_foot': 'Dieselbe Referenzantwort aus Notebook 2, Abschnitt E b. Der Nenner von Context Recall hängt davon ab, wer zerlegt.',
        'ref': 'Unternehmen können sich für Pilotprojekte bewerben und erhalten 2-3 Tage pro Woche aktive Projektbegleitung, um gemeinsam einen Prototypen zu entwickeln. Dabei können sie auch die Recheninfrastruktur nutzen.',
    },
    'en': {
        'menti1': 'Guess the number',
        'menti2': 'Guess the number, round 2',
        'code': 'Mentimeter code: ____',
        'text': 'Text',
        'defs_title': 'What is a claim?',
        'defs': [
            ('Everyday (dictionary)', 'A claim is a sentence that states something.',
             'By that the example is one claim. One could also count two: main clause and subordinate clause are two sentences.'),
            ('Logic', 'A claim is something that can be true or false on its own.',
             '"A, because B and C" are three claims: A, B, C, plus the claim that B and C are the reason for A. By that the example has three to four claims.'),
            ('NLP research (FActScore)', 'A claim is the smallest unit that carries one piece of information and can be checked against a source.',
             'Each element of a list counts. By that the example has three claims.'),
        ],
        'work_title': 'Working definition',
        'work': 'A claim is the smallest part of the reference answer that can be checked against the Kompendium on its own.',
        'work_rules': ['Every element of a list is its own claim.',
                       'Every reason (because) and every purpose (in order to) is its own claim.',
                       'Repetitions and filler do not count.'],
        'break_title': 'Example {n}: how many claims?',
        'rows': ['Everyday', 'Logic', 'Working definition'],
        'ragas_title': 'What RAGAS tells the judge',
        'ragas_lead': 'The instruction to the judge, verbatim:',
        'ragas_none': 'No definition. The judge decides for itself what a claim is.',
        'ragas_cols': [('gpt-oss-120b', '2 claims', 'mostly splits per sentence'),
                       ('llama-3-3-70b', '4 claims', 'splits finer')],
        'ragas_foot': 'The same reference answer from notebook 2, section E b. The denominator of Context Recall depends on who splits.',
        'ref': 'Companies can apply for pilot projects and receive 2 to 3 days per week of active project support to develop a prototype together. They can also use the computing infrastructure.',
    },
}
RAGAS_INSTRUCTION = ('Given a context, and an answer, analyze each statement in the answer and classify if the statement '
                     'can be attributed to the given context or not.')

# The three texts (gold 33, 17, 27) and their breakdowns: per row a list of boxes and the count shown at the right.
EXAMPLES = {
    'de': [
        ('Administrationskonten sollten durch Mehr-Faktor-Authentisierung geschützt werden, da sie ein besonders hohes Schadenspotenzial besitzen und bevorzugtes Angriffsziel darstellen.',
         [(['Administrationskonten sollten durch MFA geschützt werden, da sie ein hohes Schadenspotenzial besitzen und bevorzugtes Angriffsziel darstellen.'], '1 (oder 2)'),
          (['Adminkonten sollten durch MFA geschützt werden', 'sie besitzen ein hohes Schadenspotenzial', 'sie sind bevorzugtes Angriffsziel', 'weil: das ist der Grund'], '3 bis 4'),
          (['Adminkonten sollten durch MFA geschützt werden', 'Grund: hohes Schadenspotenzial', 'Grund: bevorzugtes Angriffsziel'], '3')]),
        ('Erstellen Sie eine Cloud-Nutzungsstrategie (A1), eine Sicherheitsrichtlinie (A2), Service-Definitionen (A3) und legen Sie Verantwortungsbereiche sowie Schnittstellen fest (A4). Diese vier Basis-Anforderungen müssen vorrangig erfüllt werden.',
         [(['Erstellen Sie ... (A4).', 'Diese vier Basis-Anforderungen müssen vorrangig erfüllt werden.'], '2'),
          (['Strategie (A1)', 'Richtlinie (A2)', 'Service-Definitionen (A3)', 'Zuständigkeiten (A4)', 'vorrangig erfüllen', 'es sind vier'], '5 bis 6'),
          (['Strategie (A1)', 'Richtlinie (A2)', 'Service-Definitionen (A3)', 'Zuständigkeiten (A4)', 'vorrangig erfüllen'], '5')]),
        ('Bei personellen Veränderungen müssen nicht mehr benötigte Benutzendenkennungen und Berechtigungen unverzüglich entfernt werden, um unberechtigte Zugriffe ehemaliger oder versetzter Mitarbeitender zu vermeiden. Es dürfen Berechtigungen nur hinsichtlich des tatsächlichen Bedarfs zur Aufgabenerfüllung vergeben werden.',
         [(['Bei personellen Veränderungen ... zu vermeiden.', 'Es dürfen Berechtigungen nur ... vergeben werden.'], '2'),
          (['Kennungen entfernen', 'Berechtigungen entfernen', 'unverzüglich', 'Zweck: unberechtigte Zugriffe vermeiden', 'Berechtigungen nur nach Bedarf'], '4 bis 5'),
          (['Kennungen und Berechtigungen entfernen', 'unverzüglich', 'Zweck: unberechtigte Zugriffe vermeiden', 'Berechtigungen nur nach Bedarf'], '4')]),
    ],
    'en': [
        ('Administrative accounts should be protected by multi-factor authentication, because they have a particularly high potential for damage and represent a preferred attack target.',
         [(['Administrative accounts should be protected by MFA, because they have a high potential for damage and are a preferred attack target.'], '1 (or 2)'),
          (['admin accounts should be protected by MFA', 'they have a high potential for damage', 'they are a preferred attack target', 'because: that is the reason'], '3 to 4'),
          (['admin accounts should be protected by MFA', 'reason: high potential for damage', 'reason: preferred attack target'], '3')]),
        ('Create a cloud usage strategy (A1), a security policy (A2), service definitions (A3), and define responsibility areas as well as interfaces (A4). These four basic requirements must be met as a priority.',
         [(['Create ... (A4).', 'These four basic requirements must be met as a priority.'], '2'),
          (['cloud usage strategy (A1)', 'security policy (A2)', 'service definitions (A3)', 'responsibilities (A4)', 'meet as a priority', 'there are four'], '5 to 6'),
          (['cloud usage strategy (A1)', 'security policy (A2)', 'service definitions (A3)', 'responsibilities (A4)', 'meet as a priority'], '5')]),
        ('When personnel changes occur, user IDs and permissions that are no longer needed must be removed promptly to avoid unauthorized access by former or transferred employees. Permissions may only be granted based on the actual need for task fulfillment.',
         [(['When personnel changes occur ... employees.', 'Permissions may only be granted ... fulfillment.'], '2'),
          (['remove user IDs', 'remove permissions', 'promptly', 'purpose: avoid unauthorized access', 'permissions only as needed'], '4 to 5'),
          (['remove user IDs and permissions', 'promptly', 'purpose: avoid unauthorized access', 'permissions only as needed'], '4')]),
    ],
}


def clear_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def new_slide(prs, title):
    layout = next(l for l in prs.slide_layouts if l.name == 'Inhalt')
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type == 1:  # TITLE
            ph.text = title
        else:
            ph._element.getparent().remove(ph._element)
    return s


def text(slide, x, y, w, h, s, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Inches(0.12))
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    lines = s if isinstance(s, list) else [s]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return shp


def menti_slide(prs, L, title, texts, with_definition):
    s = new_slide(prs, title)
    y = 1.3
    if with_definition:
        text(s, 0.77, y, 11.8, 0.95, L['work'], size=15, fill=NOTE, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.15
    h = (6.2 - y) / 3 - 0.15
    for i, (txt, _) in enumerate(texts):
        text(s, 0.77, y, 1.2, h, f"{L['text']} {i + 1}", size=14, color=GREY, bold=True)
        text(s, 2.0, y, 10.6, h, txt, size=14 if with_definition else 15, line=RULE, anchor=MSO_ANCHOR.MIDDLE)
        y += h + 0.15
    text(s, 8.6, 6.35, 4.0, 0.45, L['code'], size=14, color=ACCENT, bold=True, align=PP_ALIGN.RIGHT)
    return s


def defs_slide(prs, L):
    s = new_slide(prs, L['defs_title'])
    w = 3.75
    for i, (head, defn, note) in enumerate(L['defs']):
        x = 0.77 + i * (w + 0.3)
        text(s, x, 1.4, w, 0.6, head, size=18, bold=True)
        text(s, x, 2.05, w, 1.8, defn, size=16, line=RULE, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x, 4.0, w, 2.2, note, size=13, color=GREY)
    return s


def work_slide(prs, L):
    s = new_slide(prs, L['work_title'])
    text(s, 0.77, 1.5, 11.8, 1.3, L['work'], size=22, fill=NOTE, anchor=MSO_ANCHOR.MIDDLE)
    for i, rule in enumerate(L['work_rules']):
        text(s, 0.77, 3.2 + i * 0.8, 0.5, 0.6, str(i + 1), size=18, color=GREY, bold=True)
        text(s, 1.3, 3.2 + i * 0.8, 11.2, 0.6, rule, size=16, anchor=MSO_ANCHOR.MIDDLE)
    return s


def breakdown_slide(prs, L, n, sentence, rows):
    s = new_slide(prs, L['break_title'].format(n=n))
    text(s, 0.77, 1.25, 11.8, 1.05, sentence, size=14, line=RULE, anchor=MSO_ANCHOR.MIDDLE)
    y = 2.55
    row_h = 1.15
    for label, (boxes, count) in zip(L['rows'], rows):
        text(s, 0.77, y, 2.0, row_h, label, size=12, color=GREY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        x, avail = 2.75, 8.4
        gap = 0.1
        bw = (avail - gap * (len(boxes) - 1)) / len(boxes)
        for b in boxes:
            text(s, x, y + 0.12, bw, row_h - 0.24, b, size=10 if len(boxes) > 4 else 11 if len(boxes) > 3 else 12, line=INK, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            x += bw + gap
        text(s, 11.3, y, 1.3, row_h, count, size=24, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y += row_h + 0.12
    return s


def ragas_slide(prs, L):
    s = new_slide(prs, L['ragas_title'])
    text(s, 0.77, 1.3, 11.8, 0.4, L['ragas_lead'], size=14, color=GREY)
    text(s, 0.77, 1.75, 11.8, 1.0, RAGAS_INSTRUCTION, size=15, line=RULE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.77, 2.9, 11.8, 0.5, L['ragas_none'], size=16, bold=True)
    text(s, 0.77, 3.5, 11.8, 0.9, L['ref'], size=13, color=GREY, line=RULE, anchor=MSO_ANCHOR.MIDDLE)
    for i, (model, count, how) in enumerate(L['ragas_cols']):
        x = 0.77 + i * 6.0
        text(s, x, 4.4, 5.8, 0.45, model, size=16, bold=True)
        text(s, x, 4.85, 5.8, 0.65, count, size=28, bold=True)
        text(s, x, 5.45, 5.8, 0.4, how, size=13, color=GREY)
    text(s, 0.77, 5.85, 11.8, 0.4, L['ragas_foot'], size=12, color=GREY)
    return s


def build(lang, out):
    L, ex = T[lang], EXAMPLES[lang]
    prs = Presentation(TEMPLATE)
    clear_slides(prs)
    menti_slide(prs, L, L['menti1'], ex, with_definition=False)
    defs_slide(prs, L)
    work_slide(prs, L)
    for n, (sentence, rows) in enumerate(ex, start=1):
        breakdown_slide(prs, L, n, sentence, rows)
    menti_slide(prs, L, L['menti2'], ex, with_definition=True)
    ragas_slide(prs, L)
    prs.save(out)
    print(out, len(prs.slides), 'slides')


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
